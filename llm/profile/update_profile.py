#!/usr/bin/env python3
"""
profile フォルダの PDF / Excel からプロフィールを作り、記憶に反映する

- profile/ に置いた .pdf .xlsx .xlsm を読み、ローカルLLMで要点にまとめて profile.json に書き込む
- 追加・更新・削除されたファイルだけを処理（中身のハッシュで判定するので、無駄に作り直さない）
- profile.json には最新の資料から作った行だけを残す
  （古い行・手で書いた行は profile_backup.json に日時付きで追記して保管。毎回ファイルは増えない）
- 反映後、記憶（data/llm_memory.json）・LM Studio用テキスト・Ollamaモデルも自動で更新
- chat_memory.py の起動時にも自動でチェックされる

必要なライブラリ（初回のみ）:
    pip3 install pypdf openpyxl

使い方:
    python3 ~/MyTool/llm/profile/update_profile.py           # 変更があれば更新
    python3 ~/MyTool/llm/profile/update_profile.py --watch   # 見張り続けて、変化したらすぐ更新（Ctrl+Cで終了）
    python3 ~/MyTool/llm/profile/update_profile.py --force   # 全ファイルを作り直す
"""
import argparse
import hashlib
import shutil
import subprocess
import re
import sys
import time
from datetime import datetime
from pathlib import Path
from typing import Optional

PROFILE_DIR = Path(__file__).resolve().parent
sys.path.insert(0, str(PROFILE_DIR.parent))   # llm フォルダの chat_memory.py を読み込む
sys.path.insert(0, str(PROFILE_DIR))

from chat_memory import (BACKENDS, DATA_DIR, PROFILE_FILE, http_json, load_settings,  # noqa: E402
                         read_json, resolve_model, strip_think, update_json)

SOURCES_FILE = DATA_DIR / "profile_sources.json"   # 取り込み済みファイルのハッシュ
BACKUP_FILE = PROFILE_DIR / "profile_backup.json"  # 外した古いプロフィールを蓄積する1ファイル
TARGET_EXT = {".pdf", ".xlsx", ".xlsm"}
MAX_CHARS = 12000      # LLMに渡す最大文字数（1ファイルあたり）
MAX_LINES = 12         # 1ファイルから作るプロフィールの最大行数
WATCH_INTERVAL = 5     # --watch の確認間隔（秒）
EXCLUDE_HEADINGS = ("【現職】",)                       # 作らない見出し
EXCLUDE_PATTERN = re.compile(r"[＋+、・]?現職(\[[^\]]*\])?")   # 他の行に混ざった現職の記述も消す
ELAPSED_PATTERN = re.compile(r"\d+年目")                          # 時間とともに変わる「〇年目」は消す
REF_PATTERN = re.compile(r"\[No[^\]]*\]")                       # 資料内の参照番号 [No1] などは消す

PROMPT = """以下はユーザー本人のプロフィール資料（{name}）から取り出したテキストです。
チャットAIに「ユーザーについての前提知識」として渡すため、本人に関する重要な事実だけを抜き出してください。

ルール:
- 日本語で最大{max_lines}行。1行に1つの話題
- 各行は「【見出し】内容」の形式（例: 【経歴】IT業界12年目。それ以前は飲食業で調理13年）
- 資料に書かれていないことは書かない。推測しない
- 現在の勤務先・所属・現職に関することは書かない（【現職】の行は作らず、他の行にも含めない）
- 現在も続いている経験の経過年数（「12年目」「IT歴12年」など時間とともに変わる数字）は書かず、開始年（例: 2015年〜）で表す。終わった経歴の期間はそのまま書いてよい
- 資料内の参照番号（[No1] や [No4〜18] など）は書かない
- 同じ内容の英語版などは重複させない
- 行以外の前置き・説明・まとめは出力しない

--- 資料 ---
{text}
"""


# ---------- ファイルの読み取り ----------

def extract_pdf(path: Path) -> str:
    try:
        from pypdf import PdfReader
        return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)
    except ImportError:
        if shutil.which("pdftotext"):
            r = subprocess.run(["pdftotext", "-layout", str(path), "-"], capture_output=True, text=True)
            return r.stdout
        raise RuntimeError("PDFを読むには `pip3 install pypdf` を実行してください")


def extract_excel(path: Path) -> str:
    try:
        from openpyxl import load_workbook
    except ImportError:
        raise RuntimeError("Excelを読むには `pip3 install openpyxl` を実行してください")
    wb = load_workbook(str(path), read_only=True, data_only=True)
    lines = []
    for ws in wb.worksheets:
        lines.append(f"# シート: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None and str(c).strip()]
            if cells:
                lines.append(" | ".join(cells))
    wb.close()
    return "\n".join(lines)


def extract_text(path: Path) -> str:
    text = extract_pdf(path) if path.suffix.lower() == ".pdf" else extract_excel(path)
    lines = [" ".join(l.split()) for l in text.splitlines()]   # 余分な空白を詰める
    return "\n".join(l for l in lines if l)


# ---------- LLMで要約 ----------

def ask_llm(prompt: str, backend: str, model: str) -> str:
    if backend == "ollama":
        # Ollama は標準の文脈長が短いので、ネイティブAPIで長めに指定する
        base = BACKENDS["ollama"].removesuffix("/v1")
        res = http_json(f"{base}/api/chat", {
            "model": model, "stream": False, "options": {"num_ctx": 16384},
            "messages": [{"role": "user", "content": prompt}],
        })
        return res["message"]["content"]
    res = http_json(f"{BACKENDS[backend]}/chat/completions", {
        "model": model, "stream": False, "messages": [{"role": "user", "content": prompt}],
    })
    return res["choices"][0]["message"]["content"]


def to_profile_lines(answer: str) -> list[str]:
    lines = []
    for line in strip_think(answer).splitlines():
        line = line.strip().lstrip("-・*•0123456789.） )").strip()
        if line.startswith(EXCLUDE_HEADINGS):
            continue
        line = REF_PATTERN.sub("", ELAPSED_PATTERN.sub("", EXCLUDE_PATTERN.sub("", line)))
        if len(line) > 3 and line not in lines:
            lines.append(line)
    tagged = [l for l in lines if l.startswith("【")]   # 形式どおりの行があれば、それだけ使う
    return (tagged or lines)[:MAX_LINES]


def summarize_file(path: Path, backend: str, model: str) -> list[str]:
    text = extract_text(path)
    if not text:
        raise RuntimeError("文字を取り出せませんでした（画像だけのPDFの可能性があります）")
    if len(text) > MAX_CHARS:
        print(f"  （長いので先頭 {MAX_CHARS} 文字だけ使います）")
        text = text[:MAX_CHARS]
    lines = to_profile_lines(ask_llm(
        PROMPT.format(name=path.name, max_lines=MAX_LINES, text=text), backend, model))
    if not lines:
        raise RuntimeError("LLMの回答からプロフィールを作れませんでした")
    return lines


# ---------- 変更の検出と反映 ----------

def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def target_files() -> list[Path]:
    return sorted(p for p in PROFILE_DIR.iterdir()
                  if p.is_file() and p.suffix.lower() in TARGET_EXT
                  and not p.name.startswith(("~$", ".")))   # Excelの一時ファイル等は除外


def find_changes(force=False):
    known = read_json(SOURCES_FILE, {})
    current = {p.name: p for p in target_files()}
    changed = [p for name, p in current.items()
               if force or known.get(name, {}).get("sha256") != sha256(p)]
    deleted = [name for name in known if name not in current]
    return changed, deleted


def archive(facts: list[dict], reason: str) -> None:
    """外したプロフィール行をバックアップファイル1つに追記する"""
    if not facts:
        return
    entry = {"archived_at": datetime.now().strftime("%Y-%m-%d %H:%M"), "reason": reason,
             "facts": [{k: f[k] for k in ("text", "date", "source_file") if k in f} for f in facts]}
    update_json(BACKUP_FILE, {"history": []}, lambda b: b["history"].append(entry))
    print(f"  古いプロフィール {len(facts)}行を {BACKUP_FILE.name} に保管しました")


def check_and_update(backend: Optional[str] = None, model: Optional[str] = None,
                     force=False, sync=True) -> bool:
    """変更があれば profile.json と記憶を更新。更新したら True"""
    changed, deleted = find_changes(force)
    if not changed and not deleted:
        return False
    print("[プロフィール] 資料の変化を検出: " +
          ", ".join([f"更新 {p.name}" for p in changed] + [f"削除 {n}" for n in deleted]))

    results = {}
    if changed:
        backend = backend or load_settings()["backend"]
        model = resolve_model(backend, model)
        for p in changed:
            print(f"  {p.name} を読み取り中…（{backend} / {model}）")
            try:
                results[p.name] = (summarize_file(p, backend, model), sha256(p))
            except Exception as e:   # 1ファイル失敗しても他は続ける。次回また再試行される
                print(f"  [スキップ] {p.name}: {e}")
    if not results and not deleted:
        return False

    today = datetime.now().strftime("%Y-%m-%d")
    removed = []

    def apply_profile_json(data):
        # 最新の資料から作った行だけを残す
        #   ・今回読み直したファイル → 新しい行に差し替え
        #   ・変化のないファイル（読み取り失敗を含む） → そのまま
        #   ・削除されたファイル由来の行・手書きの行 → 外してバックアップへ
        done = set(results) | set(deleted)
        alive = {p.name for p in target_files()}
        old = data.get("facts", [])
        facts = [f for f in old if f.get("source_file") in alive and f.get("source_file") not in done]
        for name, (lines, _) in results.items():
            facts += [{"text": t, "date": today, "source_file": name} for t in lines]
        new_texts = {f["text"] for f in facts}
        removed.extend(f for f in old if f["text"] not in new_texts)
        data["facts"] = facts

    def apply_sources(src):
        for name in deleted:
            src.pop(name, None)
        for name, (_, digest) in results.items():
            src[name] = {"sha256": digest, "updated": today}

    update_json(PROFILE_FILE, {"facts": []}, apply_profile_json)   # 書く直前に読み直す
    update_json(SOURCES_FILE, {}, apply_sources)
    archive(removed, "資料の更新: " + ", ".join(list(results) + [f"{n}(削除)" for n in deleted]))
    for name, (lines, _) in results.items():
        print(f"  {name} から {len(lines)}行:")
        for t in lines:
            print(f"    + {t}")

    from add_profile import apply_profile
    apply_profile(sync=sync)
    return True


def main():
    p = argparse.ArgumentParser(description="profile フォルダの PDF / Excel からプロフィールを更新")
    p.add_argument("--watch", action="store_true", help="見張り続けて、変化したらすぐ更新")
    p.add_argument("--force", action="store_true", help="全ファイルを作り直す")
    p.add_argument("--backend", choices=BACKENDS, help="要約に使うバックエンド（省略時は前回と同じ）")
    p.add_argument("--model", help="要約に使うモデル（省略時は前回と同じ）")
    a = p.parse_args()

    if not a.watch:
        if not check_and_update(a.backend, a.model, a.force):
            print("[プロフィール] 資料に変化はありません")
        return

    print(f"[プロフィール] {PROFILE_DIR} を見張っています（Ctrl+Cで終了）")
    check_and_update(a.backend, a.model, a.force)
    last = None
    try:
        while True:
            # まず更新日時とサイズだけを見る（軽い）。変わったときだけ中身を比べる
            snap = tuple((f.name, f.stat().st_mtime, f.stat().st_size) for f in target_files())
            if snap != last:
                if last is not None:
                    time.sleep(2)   # 保存が終わるのを少し待つ
                    check_and_update(a.backend, a.model)
                last = snap
            time.sleep(WATCH_INTERVAL)
    except KeyboardInterrupt:
        print("\n終了しました")


if __name__ == "__main__":
    main()
